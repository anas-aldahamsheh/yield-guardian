"""Read an uploaded table (CSV or XLSX) into a header and rows, with plain-language errors.

Both file parsers (sensor units, wafer dies) build on this: it only finds the header and the
cells, and leaves every meaning to them. Row numbers are the file's own (the header is the
first non-empty row; in a CSV it is usually row 1), so a message can say "row 14, column S060".
"""

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import PurePath

from openpyxl import load_workbook

MAX_BYTES = 32 * 1024 * 1024  # 5,000 SECOM units as CSV are about 24 MB
MAX_XLSX_UNPACKED_BYTES = 400 * 1024 * 1024
MAX_ROWS = 2_000_000
MAX_COLUMNS = 1024
CSV_SUFFIXES = (".csv", ".txt")
XLSX_SUFFIXES = (".xlsx", ".xlsm")

Cell = str | float | int | bool | datetime | None


class TableError(Exception):
    """The file cannot be read as a table; `message` says why and what to do."""

    def __init__(self, message: str, *, row: int | None = None, column: str | None = None) -> None:
        super().__init__(message)
        self.message = message
        self.row = row
        self.column = column


@dataclass(frozen=True, slots=True)
class Row:
    number: int  # 1-based row in the file
    cells: list[Cell]  # aligned with Table.header; short rows are padded with None


@dataclass(frozen=True, slots=True)
class Table:
    header: list[str]
    header_row: int
    rows: list[Row]
    kind: str  # "csv" or "xlsx"


def read_table(data: bytes, filename: str) -> Table:
    """The header and the non-empty rows of the first sheet (XLSX) or the file (CSV)."""
    if len(data) == 0:
        raise TableError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise TableError(
            f"The file is {len(data) / 1_048_576:.1f} MB; the limit is "
            f"{MAX_BYTES // 1_048_576} MB. Split it into smaller files."
        )
    suffix = PurePath(filename).suffix.lower()
    if suffix in CSV_SUFFIXES:
        return _from_records(_csv_records(data), "csv")
    if suffix in XLSX_SUFFIXES:
        return _from_records(_xlsx_records(data), "xlsx")
    raise TableError(f'"{suffix or filename}" files are not read here. Use .csv or .xlsx.')


def _csv_records(data: bytes) -> list[tuple[int, list[Cell]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise TableError('The CSV is not UTF-8 text. Save it as "CSV UTF-8".') from error
    sample = text[:65536]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    records: list[tuple[int, list[Cell]]] = []
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    try:
        for number, values in enumerate(reader, start=1):
            if number > MAX_ROWS + 1:
                raise TableError(f"The file has more than {MAX_ROWS:,} rows.")
            records.append((number, list(values)))
    except csv.Error as error:
        raise TableError(f"The CSV cannot be read: {error}.", row=reader.line_num) from error
    return records


def _xlsx_records(data: bytes) -> list[tuple[int, list[Cell]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
    except Exception as error:  # BadZipFile, and NotImplementedError for a damaged version field
        raise TableError("The file is not a valid .xlsx workbook.") from error
    if unpacked > MAX_XLSX_UNPACKED_BYTES:
        raise TableError("The workbook is too large once unpacked. Save the sheet as CSV.")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many types for damaged files
        raise TableError("The workbook cannot be opened. Save the sheet as CSV.") from error
    try:
        if not workbook.worksheets:
            raise TableError("The workbook has no sheet.")
        sheet = workbook.worksheets[0]
        records: list[tuple[int, list[Cell]]] = []
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if number > MAX_ROWS + 1:
                raise TableError(f"The sheet has more than {MAX_ROWS:,} rows.")
            records.append((number, [_xlsx_cell(v) for v in values]))
        return records
    except TableError:
        raise
    except Exception as error:  # read-only sheets are parsed lazily: damage shows up here
        raise TableError(
            "The workbook's first sheet is damaged and cannot be read. Save the sheet as CSV."
        ) from error
    finally:
        workbook.close()


def _xlsx_cell(value: object) -> Cell:
    if value is None or isinstance(value, str | bool | int | float | datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time())
    if isinstance(value, time):
        return value.isoformat()
    return str(value)


def _blank(cell: Cell) -> bool:
    return cell is None or (isinstance(cell, str) and cell.strip() == "")


def _from_records(records: list[tuple[int, list[Cell]]], kind: str) -> Table:
    non_empty = [(n, cells) for n, cells in records if not all(_blank(c) for c in cells)]
    if not non_empty:
        raise TableError("The file has no rows.")
    header_row, raw_header = non_empty[0]
    while raw_header and _blank(raw_header[-1]):
        raw_header = raw_header[:-1]
    if len(raw_header) > MAX_COLUMNS:
        raise TableError(f"The header has {len(raw_header)} columns; the limit is {MAX_COLUMNS}.")
    header: list[str] = []
    for position, cell in enumerate(raw_header, start=1):
        name = "" if cell is None else str(cell).strip()
        if not name:
            raise TableError(f"Column {position} of the header is empty.", row=header_row)
        if name in header:
            raise TableError(f'The header names "{name}" twice.', row=header_row, column=name)
        header.append(name)
    width = len(header)
    rows: list[Row] = []
    for number, cells in non_empty[1:]:
        extra = cells[width:]
        if any(not _blank(c) for c in extra):
            raise TableError(
                f"Row {number} has more values than the header has columns ({width}).",
                row=number,
            )
        padded = list(cells[:width]) + [None] * max(0, width - len(cells))
        rows.append(Row(number=number, cells=[None if _blank(c) else c for c in padded]))
    return Table(header=header, header_row=header_row, rows=rows, kind=kind)
