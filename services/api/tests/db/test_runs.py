"""Live runs in PostgreSQL: the store (items, stages, unit charts, events, expiry, retention,
uploads) and the run and upload routes end to end, with the real SECOM engine and sensor-file
parser on real SECOM units and files. Wafer tracks use the wafer double here (fast); the real
wafer engine runs end to end in tests/db/test_runs_real.py."""

import asyncio
import base64
import io
import re
from collections.abc import AsyncIterator
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from random import Random

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from tests.db.conftest import clean
from tests.db.fakes import FakeRunner
from tests.live.doubles import (
    SECOM_DATA,
    WAFERMAP_DEMO,
    WaferDouble,
    demo_wafers,
    held_out_ids,
    real_secom,
    real_sensor_engine,
)
from yield_guardian.api.app import create_app
from yield_guardian.contracts.runs import (
    RunAlarm,
    RunOrigin,
    SensorUnitInput,
    StageStartedEvent,
    UnitChartResult,
    UploadSummary,
)
from yield_guardian.contracts.services import NotFoundError
from yield_guardian.core.settings import Settings
from yield_guardian.db.tables import InvestigationRow
from yield_guardian.live.runs import Engines, RunOrchestrator
from yield_guardian.live.sensor_file import SecomSensorFileParser
from yield_guardian.live.uploads import SENSOR_HEADER, WAFER_HEADER, Attribution
from yield_guardian.live.wiring import install_runs
from yield_guardian.services.persistence import Persistence, build_persistence
from yield_guardian.services.registry import ServiceRegistry
from yield_guardian.services.runs import (
    RUN_TTL,
    Z_IN_RESULT,
    DbRunStore,
    GoneError,
    RunInput,
    as_stored,
    initial_stage,
    run_done,
)

ORIGIN = RunOrigin(source="pick", datasets=["secom"], title="Your pick of real data", rule="test")
REAL_LOT = "lot10291"


def units(n: int) -> list[SensorUnitInput]:
    return real_sensor_engine().from_secom(held_out_ids()[:n])


@dataclass
class Clock:
    now: datetime

    def __call__(self) -> datetime:
        return self.now


def store_at(persistence: Persistence, clock: Clock) -> DbRunStore:
    return DbRunStore(persistence.database.sessionmaker, persistence.notifier, now=clock)


async def new_run(store: DbRunStore, data: RunInput) -> str:
    run = await store.create(ORIGIN, data, RunOrchestrator.plan(data))
    return run.run_id


def started(stage: str):  # type: ignore[no-untyped-def]
    return lambda seq, at: StageStartedEvent(seq=seq, at=at, stage=stage)


# Store ----------------------------------------------------------------------------------------


async def test_a_run_keeps_its_input_and_reads_one_item_at_a_time(
    persistence: Persistence,
) -> None:
    store = persistence.runs
    data = RunInput(units=units(25), wafers=demo_wafers(REAL_LOT)[:3])
    run_id = await new_run(store, data)
    assert re.fullmatch(r"run_[0-9a-z]{26}", run_id)
    run = await store.get(run_id)
    assert run.status == "queued"
    assert (run.units, run.wafers) == (25, 3)
    assert run.expires_at is not None
    assert run.expires_at - run.created_at == RUN_TTL
    assert [s.status for s in run.stages] == ["pending"] * 8
    assert await store.get_input(run_id) == data
    assert await store.get_unit_input(run_id, 7) == data.units[7]
    assert await store.get_wafer_input(run_id, 2) == data.wafers[2]
    assert await store.get_unit_input(run_id, 25) is None
    assert await store.get_wafer_input(run_id, -1) is None
    assert await store.get_run("run_01jaf3k9x2m4q7r8s0t1v2w3x4") is None
    again = await store.get_run(run_id)
    # Each read carries the moment it was made; the run read is the same.
    assert again is not None
    assert again.read_at is not None
    assert run.read_at is not None
    assert again.read_at >= run.read_at
    assert again.model_copy(update={"read_at": run.read_at}) == run


async def test_events_are_numbered_resume_after_a_seq_and_end_after_run_done(
    persistence: Persistence,
) -> None:
    store = persistence.runs
    run_id = await new_run(store, RunInput(units=units(3)))
    seen: list[int] = []

    async def follow() -> None:
        async for event in store.events(run_id):
            seen.append(event.seq)

    reader = asyncio.create_task(follow())
    await asyncio.sleep(0.2)
    assert seen == [1]  # run_started, then it waits for more
    await store.update(run_id, status="running", events=[started("intake")])
    await store.update(run_id, status="done", events=[run_done("done", 0, 0)])
    await asyncio.wait_for(reader, 5)  # woken by NOTIFY, closed after run_done
    assert seen == [1, 2, 3]
    resumed = [e async for e in store.events(run_id, after_seq=1)]
    assert [(e.seq, e.type) for e in resumed] == [(2, "stage_started"), (3, "run_done")]
    assert [e async for e in store.events(run_id, after_seq=3)] == []  # finished: ends at once
    run = await store.get(run_id)
    assert run.status == "done"


async def test_one_unit_chart_is_read_with_every_z(persistence: Persistence) -> None:
    store = persistence.runs
    given = units(Z_IN_RESULT + 20)
    run_id = await new_run(store, RunInput(units=given))
    with pytest.raises(NotFoundError, match="not been computed"):
        await store.unit_chart(run_id, 0)
    charts = [as_stored(u) for u in real_sensor_engine().control_charts(given).units]
    await store.update(run_id, unit_charts=charts)
    found = await store.unit_chart(run_id, Z_IN_RESULT + 5)
    assert found == charts[Z_IN_RESULT + 5]
    assert found.z is not None
    assert len(found.z) == len(real_sensor_engine().control_charts(given[:1]).features)
    assert await store.get_unit_chart(run_id, 3) == charts[3]
    with pytest.raises(NotFoundError, match="no unit 220"):
        await store.unit_chart(run_id, len(given))
    assert await store.get_unit_chart(run_id, len(given)) is None
    assert await store.get_unit_chart("run_01jaf3k9x2m4q7r8s0t1v2w3x4", 0) is None


async def test_expired_runs_are_gone_and_pruned_unless_an_investigation_uses_them(
    persistence: Persistence,
) -> None:
    clock = Clock(datetime.now(UTC))
    store = store_at(persistence, clock)
    kept = await new_run(store, RunInput(units=units(2)))
    used = await new_run(store, RunInput(units=units(2)))
    retained = await new_run(store, RunInput(units=units(2)))
    async with persistence.database.sessionmaker() as db, db.begin():
        db.add(
            InvestigationRow(
                investigation_id="inv_0123456789",
                title="from a run alarm",
                status="queued",
                trigger_kind="alarm",
                trigger_ref=f"{used}:RA-1",
                created_at=clock.now,
                updated_at=clock.now,
            )
        )
    await store.retain(retained)
    assert (await store.get(retained)).expires_at is None
    assert (await store.get(used)).expires_at is None  # an investigation uses it

    clock.now += RUN_TTL + timedelta(seconds=1)
    with pytest.raises(GoneError) as gone:
        await store.get(kept)
    assert gone.value.status == 410
    assert await store.get_run(kept) is None
    assert await store.get_unit_input(kept, 0) is None
    report = await store.prune()
    assert report.runs == 1
    with pytest.raises(GoneError):  # pruned: its ID still says when it was made
        await store.get(kept)
    assert (await store.get(used)).run_id == used
    assert (await store.get(retained)).run_id == retained
    assert await store.get_unit_input(used, 1) is not None
    assert (await store.prune()).runs == 0


async def test_uploads_are_kept_for_24_hours(persistence: Persistence) -> None:
    clock = Clock(datetime.now(UTC))
    store = store_at(persistence, clock)
    data = RunInput(units=units(4))

    def summary(upload_id: str, expires_at: datetime) -> UploadSummary:
        return UploadSummary(
            upload_id=upload_id,
            kind="sensors",
            filename="units.csv",
            units=4,
            wafers=0,
            lots=0,
            missing_readings=0,
            issues=[],
            runnable=True,
            expires_at=expires_at,
        )

    made = await store.create_upload(
        kind="sensors", filename="units.csv", data=data, summary=summary
    )
    found, items = await store.get_upload(made.upload_id)
    assert found == made
    assert items == data
    # A second past the TTL: once pruned, the upload ID's own (wall-clock) time decides "gone".
    clock.now += RUN_TTL + timedelta(seconds=1)
    with pytest.raises(GoneError):
        await store.get_upload(made.upload_id)
    assert (await store.prune()).uploads == 1
    with pytest.raises(GoneError):
        await store.get_upload(made.upload_id)


async def test_an_interrupted_or_stale_run_ends_its_stream(persistence: Persistence) -> None:
    clock = Clock(datetime.now(UTC))
    store = store_at(persistence, clock)
    data = RunInput(units=units(3))
    run_id = await new_run(store, data)
    await store.update(
        run_id,
        status="running",
        stage=initial_stage("intake").model_copy(update={"status": "running"}),
        events=[started("intake")],
    )
    clock.now += timedelta(minutes=5)
    assert await store.fail_stale(timedelta(minutes=2)) == [run_id]
    events = [e async for e in store.events(run_id, after_seq=2)]
    assert [(e.type, getattr(e, "stage", None)) for e in events][:2] == [
        ("stage_failed", "intake"),
        ("stage_skipped", "sensor_risk"),
    ]
    assert events[-1].type == "run_done"
    run = await store.get(run_id)
    assert run.status == "failed"
    assert {s.status for s in run.stages} == {"failed", "skipped"}
    assert await store.interrupt(run_id) is False  # already finished


# Routes ---------------------------------------------------------------------------------------


@dataclass
class Api:
    client: httpx.AsyncClient
    registry: ServiceRegistry
    persistence: Persistence


def wafer_double() -> WaferDouble:
    return WaferDouble(demo={REAL_LOT: demo_wafers(REAL_LOT)})


@pytest.fixture
async def api(database_url: str, db_settings: Settings) -> AsyncIterator[Api]:
    await clean(database_url)
    registry = ServiceRegistry(settings=db_settings, runner=FakeRunner())
    persistence = build_persistence(db_settings, runner_provider=lambda: registry.runner)
    persistence.install(registry)
    install_runs(
        registry,
        persistence.runs,
        secom=real_secom(),
        wafermap=None,
        secom_attribution=SECOM_DATA / "ATTRIBUTION.md",
        engines=Engines(
            sensors=real_sensor_engine(),
            wafers=wafer_double(),
            sensor_file=SecomSensorFileParser(),
        ),
    )
    assert registry.uploads is not None
    registry.uploads.attribution = Attribution.load(
        SECOM_DATA / "ATTRIBUTION.md", (WAFERMAP_DEMO / "ATTRIBUTION.md").read_text()
    )
    app = create_app(db_settings, registry)
    async with app.router.lifespan_context(app):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            yield Api(client, registry, persistence)


def parse_sse(body: str) -> list[dict[str, str]]:
    events: list[dict[str, str]] = []
    for block in re.split(r"\r?\n\r?\n", body):
        fields: dict[str, str] = {}
        for line in block.splitlines():
            if not line or line.startswith(":"):
                continue
            name, _, value = line.partition(":")
            fields[name] = value[1:] if value.startswith(" ") else value
        if "data" in fields:
            events.append(fields)
    return events


async def stream(
    client: httpx.AsyncClient, run_id: str, after: int | None = None
) -> list[dict[str, str]]:
    headers = {"Last-Event-ID": str(after)} if after is not None else {}
    async with client.stream("GET", f"/api/v1/runs/{run_id}/events", headers=headers) as response:
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/event-stream")
        body = "".join([chunk async for chunk in response.aiter_text()])
    return parse_sse(body)


async def test_a_sample_run_streams_every_stage_in_order_and_resumes(api: Api) -> None:
    created = await api.client.post("/api/v1/runs", json={"source": "sample", "seed": 11})
    assert created.status_code == 201
    body = created.json()
    run_id = body["run_id"]
    assert body["events_url"] == f"/api/v1/runs/{run_id}/events"
    assert body["origin"]["seed"] == 11
    events = await stream(api.client, run_id)
    assert [int(e["id"]) for e in events] == list(range(1, len(events) + 1))
    assert events[0]["event"] == "run_started"
    assert events[-1]["event"] == "run_done"
    stages = [
        re.search(r'"stage":"(\w+)"', e["data"]).group(1)  # type: ignore[union-attr]
        for e in events
        if e["event"] in {"stage_result", "stage_skipped", "stage_failed"}
    ]
    assert stages == [
        "intake",
        "sensor_risk",
        "control_charts",
        "drift_health",
        "wafer_pattern",
        "wafer_similar",
        "lot",
        "final_test",
    ]
    assert all(e["event"] != "stage_failed" for e in events)
    resumed = await stream(api.client, run_id, after=4)
    assert [int(e["id"]) for e in resumed] == list(range(5, len(events) + 1))
    assert [e["data"] for e in resumed] == [e["data"] for e in events[4:]]

    run = (await api.client.get(f"/api/v1/runs/{run_id}")).json()
    assert run["status"] == "done"
    assert run["units"] == 40
    charts = next(s for s in run["stages"] if s["stage"] == "control_charts")["result"]
    assert all(u["z"] is not None for u in charts["units"])  # 40 units, all within the first 200
    one = await api.client.get(f"/api/v1/runs/{run_id}/units/39/chart")
    assert one.status_code == 200
    assert UnitChartResult.model_validate(one.json()).z == charts["units"][39]["z"]
    missing = await api.client.get(f"/api/v1/runs/{run_id}/units/40/chart")
    assert missing.status_code == 404
    assert missing.headers["content-type"] == "application/problem+json"
    for alarm in run["alarms"]:
        RunAlarm.model_validate(alarm)
    again = await api.client.post("/api/v1/runs", json={"source": "sample", "seed": 11})
    assert again.json()["origin"] == body["origin"]  # the same seed, the same real items


async def test_run_problems_are_problem_json(api: Api, db_settings: Settings) -> None:
    unknown = await api.client.get("/api/v1/runs/run_01jaf3k9x2m4q7r8s0t1v2w3x4")
    assert unknown.status_code == 410  # an ID made long ago (by its ULID time) has expired
    assert unknown.json()["code"] == "gone"
    fresh = api.registry.runs
    assert fresh is not None
    run_id = await new_run(api.persistence.runs, RunInput(units=units(2)))
    unknown_now = run_id[:-6] + "zzzzzz"
    response = await api.client.get(f"/api/v1/runs/{unknown_now}")
    assert response.status_code == 404
    assert response.json()["code"] == "not_found"
    assert (await api.client.get("/api/v1/runs/not-a-run")).status_code == 422
    async with api.persistence.database.sessionmaker() as db, db.begin():
        await db.execute(
            text("UPDATE runs SET expires_at = now() - interval '1 minute' WHERE run_id = :id"),
            {"id": run_id},
        )
    for path in (f"/api/v1/runs/{run_id}", f"/api/v1/runs/{run_id}/events"):
        expired = await api.client.get(path)
        assert expired.status_code == 410, path
        assert expired.headers["content-type"] == "application/problem+json"
        assert "24 hours" in expired.json()["detail"]
    assert (await api.client.get(f"/api/v1/runs/{run_id}/units/0/chart")).status_code == 410
    bad_pick = await api.client.post("/api/v1/runs", json={"source": "pick", "lot": "nolot"})
    assert bad_pick.status_code == 404
    repeat = await api.client.post("/api/v1/runs", json={"source": "pick", "unit_ids": [3, 3]})
    assert repeat.status_code == 422
    assert repeat.json()["code"] == "validation_error"
    bad_case = await api.client.post("/api/v1/runs", json={"source": "case", "case_id": "nope"})
    assert bad_case.status_code == 404
    drawn = {
        "source": "drawn",
        "wafers": [{"lot": "mine", "wafer": 1, "map": {"height": 2, "width": 2, "cells": "AAAA"}}],
    }
    full = base64.b64encode(b"\x01" * 512 * 512).decode()  # a 512 x 512 map, built here
    huge = {
        "source": "drawn",
        "wafers": [
            {"lot": "mine", "wafer": n, "map": {"height": 512, "width": 512, "cells": full}}
            for n in range(1, 9)
        ],
    }
    too_many = await api.client.post("/api/v1/runs", json=huge)
    assert too_many.status_code == 422
    assert too_many.json()["detail"] == (
        "This run has 2,097,152 dies; the limit is 2,000,000. Choose fewer wafers."
    )
    no_die = await api.client.post("/api/v1/runs", json=drawn)
    assert no_die.status_code == 422
    assert "Drawing 1: The map has 3 cells" in no_die.json()["detail"]

    bare = create_app(db_settings, ServiceRegistry(settings=db_settings))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=bare), base_url="http://test"
    ) as client:
        for method, path in (
            ("GET", f"/api/v1/runs/{run_id}"),
            ("GET", "/api/v1/runs/cases"),
            ("POST", "/api/v1/runs"),
            ("GET", "/api/v1/uploads/examples/sensors.csv"),
        ):
            answer = await client.request(method, path, json={"source": "sample"})
            assert answer.status_code == 503, path
            assert answer.json()["code"] == "not_ready"


async def test_the_case_catalog_runs_what_its_rule_found(api: Api) -> None:
    listed = (await api.client.get("/api/v1/runs/cases")).json()["items"]
    ids = [c["case_id"] for c in listed]
    assert ids[:2] == ["spe_alarm", "slow_drift"]
    assert (await api.client.get("/api/v1/runs/cases")).json()["items"] == listed
    created = await api.client.post("/api/v1/runs", json={"source": "case", "case_id": "spe_alarm"})
    assert created.status_code == 201
    run_id = created.json()["run_id"]
    events = await stream(api.client, run_id)
    alarms = [e["data"] for e in events if e["event"] == "alarm"]
    assert any('"kind":"spe"' in a and '"severity":"alarm"' in a for a in alarms)
    layouts = await api.client.get("/api/v1/runs/layouts")
    assert layouts.status_code == 200
    assert layouts.json() == {"items": []}  # the double has none; tests/db/test_runs_real.py


# Uploads --------------------------------------------------------------------------------------


async def upload(client: httpx.AsyncClient, name: str, content: bytes, kind: str | None = None):  # type: ignore[no-untyped-def]
    data = {"kind": kind} if kind else None
    return await client.post("/api/v1/uploads", files={"file": (name, content)}, data=data)


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
async def test_a_real_example_file_uploads_back_and_runs(api: Api, fmt: str) -> None:
    example = await api.client.get(f"/api/v1/uploads/examples/sensors.{fmt}")
    assert example.status_code == 200
    assert example.headers["content-disposition"] == (
        f'attachment; filename="yield-guardian-sensors-example.{fmt}"'
    )
    assert "McCann" in example.headers["x-data-attribution"]
    assert example.headers["link"] == '<https://doi.org/10.24432/C54305>; rel="cite-as"'
    assert "test block" in example.headers["x-example-rule"]
    parsed = SecomSensorFileParser().parse(example.content, f"example.{fmt}")
    assert parsed.issues == []
    expected = real_sensor_engine().sample(Random(0)).items
    assert [u.unit_id for u in parsed.items] == [u.unit_id for u in expected]
    assert [u.readings for u in parsed.items] == [u.readings for u in expected]
    assert [u.timestamp for u in parsed.items] == [u.timestamp for u in expected]
    assert [u.result for u in parsed.items] == [u.result for u in expected]

    sent = await upload(api.client, f"mine.{fmt}", example.content)
    assert sent.status_code == 201
    summary = UploadSummary.model_validate(sent.json())
    assert summary.kind == "sensors"
    assert summary.runnable is True
    assert summary.issues == []
    assert summary.units == len(expected)
    assert summary.missing_readings == sum(r is None for u in expected for r in u.readings)
    created = await api.client.post(
        "/api/v1/runs", json={"source": "upload", "upload_id": summary.upload_id}
    )
    assert created.status_code == 201
    assert created.json()["origin"]["title"] == "Your file"
    assert created.json()["origin"]["filename"] == f"mine.{fmt}"
    events = await stream(api.client, created.json()["run_id"])
    assert events[-1]["event"] == "run_done"
    assert '"status":"done"' in events[-1]["data"]
    run = (await api.client.get(f"/api/v1/runs/{created.json()['run_id']}")).json()
    risk = next(s for s in run["stages"] if s["stage"] == "sensor_risk")["result"]
    assert {u["validity"] for u in risk["units"]} == {"new_data"}  # a file is not SECOM itself
    if fmt == "xlsx":
        about = load_workbook(io.BytesIO(example.content))["About"]
        lines = [row[0] for row in about.iter_rows(values_only=True) if row[0]]
        assert any("CC BY 4.0" in str(line) for line in lines)


async def test_templates_are_the_header_only(api: Api) -> None:
    for kind, header in (("sensors", SENSOR_HEADER), ("wafers", WAFER_HEADER)):
        csv_file = await api.client.get(f"/api/v1/uploads/templates/{kind}.csv")
        assert csv_file.status_code == 200
        assert csv_file.text.strip() == ",".join(header)
        xlsx = await api.client.get(f"/api/v1/uploads/templates/{kind}.xlsx")
        sheet = load_workbook(io.BytesIO(xlsx.content)).worksheets[0]
        assert [list(r) for r in sheet.iter_rows(values_only=True)] == [header]
    assert (await api.client.get("/api/v1/uploads/templates/other.csv")).status_code == 422
    empty = await upload(api.client, "template.csv", (",".join(SENSOR_HEADER) + "\r\n").encode())
    assert empty.status_code == 201  # readable: what is wrong comes back as issues
    summary = empty.json()
    assert summary["runnable"] is False
    assert summary["issues"][0]["severity"] == "error"
    refused = await api.client.post(
        "/api/v1/runs", json={"source": "upload", "upload_id": summary["upload_id"]}
    )
    assert refused.status_code == 422
    assert "Fix them and upload it again" in refused.json()["detail"]


async def test_upload_problems(api: Api) -> None:
    for name, content in (
        ("notes.pdf", b"%PDF-1.7"),
        ("broken.xlsx", b"PK\x03\x04 not a workbook"),
        ("latin.csv", "unit_id;S001\nä;1".encode("latin-1")),
        ("plain.csv", b"name,colour\nx,red\n"),
    ):
        answer = await upload(api.client, name, content)
        assert answer.status_code == 422, name
        assert answer.json()["code"] == "validation_error"
    wrong = await upload(
        api.client, "units.csv", b"lot,wafer,x,y,result\nA,1,0,0,pass\n", "sensors"
    )
    assert wrong.status_code == 201
    assert wrong.json()["runnable"] is False
    no_parser = await upload(api.client, "dies.csv", b"lot,wafer,x,y,result\nA,1,0,0,pass\n")
    assert no_parser.status_code == 503  # the double setup has no wafer-file parser
    too_big = await api.client.post(
        "/api/v1/uploads",
        files={"file": ("big.csv", b"0" * (33 * 1024 * 1024))},
    )
    assert too_big.status_code == 413
    assert too_big.json()["code"] == "bad_request"
    assert too_big.json()["detail"] == "The file is larger than 32 MB. Split it into smaller files."
    gone = await api.client.post(
        "/api/v1/runs", json={"source": "upload", "upload_id": "upl_01jaf3k9x2m4q7r8s0t1v2w3x4"}
    )
    assert gone.status_code == 410


async def test_the_wafer_example_is_every_die_of_a_real_lot(api: Api) -> None:
    example = await api.client.get("/api/v1/uploads/examples/wafers.csv")
    assert example.status_code == 200
    assert "Wu, M.-J." in example.headers["x-data-attribution"]
    assert example.headers["link"] == '<https://doi.org/10.1109/TSM.2014.2364237>; rel="cite-as"'
    lines = example.text.strip().splitlines()
    assert lines[0] == ",".join(WAFER_HEADER)
    real = demo_wafers(REAL_LOT)
    grids: dict[int, dict[tuple[int, int], int]] = {}
    for line in lines[1:]:
        lot, number, x, y, result = line.split(",")
        assert lot == REAL_LOT
        grids.setdefault(int(number), {})[(int(x), int(y))] = 1 if result == "pass" else 2
    assert sorted(grids) == [w.wafer for w in real]
    for wafer in real:
        cells = base64.b64decode(wafer.map.cells)
        rebuilt = bytearray(wafer.map.height * wafer.map.width)
        for (col, row), value in grids[wafer.wafer].items():
            rebuilt[row * wafer.map.width + col] = value
        assert bytes(rebuilt) == cells
