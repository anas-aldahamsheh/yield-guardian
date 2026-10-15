"""Live runs end to end with both real engines (V2 sensors, V3 wafers) on the real bundles:
the API wired as bootstrap wires it, a sample run, real example files uploaded back, a drawing
on a real die layout, and the case catalog computed from the data."""

import asyncio
import base64
import functools
import io
import time
from collections.abc import AsyncIterator
from pathlib import Path
from random import Random

import httpx
import pytest
from openpyxl import load_workbook

from tests.db.conftest import clean
from tests.db.fakes import FakeRunner
from tests.db.test_runs import parse_sse
from tests.live.doubles import ARTIFACTS, SECOM_DATA, real_secom, real_sensor_engine
from yield_guardian.api.app import create_app
from yield_guardian.bundles.manifest import latest_version, verify_bundle
from yield_guardian.bundles.wafermap import WafermapBundle, load_wafermap_bundle
from yield_guardian.contracts.runs import (
    LotResult,
    Run,
    RunEvent,
    RunEventPayload,
    StageResultEvent,
    WaferPatternResult,
)
from yield_guardian.contracts.wafer import WaferMap
from yield_guardian.core.settings import Settings
from yield_guardian.live.causes import LiteratureCauses
from yield_guardian.live.runs import Engines
from yield_guardian.live.sensor_file import SecomSensorFileParser
from yield_guardian.live.wafer_file import WaferDieFile
from yield_guardian.live.wafers import NEEDS_LOT, LiveWafers, decode_map
from yield_guardian.live.wiring import install_runs
from yield_guardian.services.persistence import build_persistence
from yield_guardian.services.registry import ServiceRegistry

KNOWLEDGE = Path(__file__).resolve().parents[4] / "docs" / "knowledge"


@functools.cache
def real_wafermap() -> WafermapBundle:
    return load_wafermap_bundle(verify_bundle(latest_version(ARTIFACTS, "wafermap"), "wafermap"))


@functools.cache
def real_wafer_engine() -> LiveWafers:
    return LiveWafers(real_wafermap(), LiteratureCauses.from_directory(KNOWLEDGE))


@pytest.fixture
async def api(database_url: str, db_settings: Settings) -> AsyncIterator[httpx.AsyncClient]:
    await clean(database_url)
    registry = ServiceRegistry(settings=db_settings, runner=FakeRunner())
    persistence = build_persistence(db_settings, runner_provider=lambda: registry.runner)
    persistence.install(registry)
    install_runs(
        registry,
        persistence.runs,
        secom=real_secom(),
        wafermap=real_wafermap(),
        secom_attribution=SECOM_DATA / "ATTRIBUTION.md",
        engines=Engines(
            sensors=real_sensor_engine(),
            wafers=real_wafer_engine(),
            sensor_file=SecomSensorFileParser(),
            wafer_file=WaferDieFile(),
        ),
    )
    app = create_app(db_settings, registry)
    async with app.router.lifespan_context(app):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport, base_url="http://test", timeout=120
        ) as client:
            yield client


async def run_to_end(
    client: httpx.AsyncClient, body: dict[str, object]
) -> tuple[Run, list[RunEventPayload]]:
    created = await client.post("/api/v1/runs", json=body)
    assert created.status_code == 201, created.text
    run_id = created.json()["run_id"]
    async with client.stream("GET", f"/api/v1/runs/{run_id}/events") as response:
        text = "".join([chunk async for chunk in response.aiter_text()])
    events = [RunEvent.model_validate_json(e["data"]).root for e in parse_sse(text)]
    run = Run.model_validate((await client.get(f"/api/v1/runs/{run_id}")).json())
    return run, events


def closing(events: list[RunEventPayload]) -> dict[str, str]:
    return {
        e.stage: e.type
        for e in events
        if e.type in {"stage_result", "stage_skipped", "stage_failed"} and hasattr(e, "stage")
    }


async def test_a_sample_runs_every_stage_on_both_real_engines(api: httpx.AsyncClient) -> None:
    started = time.perf_counter()
    run, events = await run_to_end(api, {"source": "sample", "seed": 2026})
    took = time.perf_counter() - started
    assert run.status == "done"
    assert run.origin.datasets == ["secom", "wm811k"]
    assert run.units == 40
    assert run.wafers >= 2
    assert set(closing(events).values()) == {"stage_result"}
    timings = {e.stage: round(e.elapsed_ms) for e in events if isinstance(e, StageResultEvent)}
    print(f"sample run seed 2026: {took:.1f} s wall, stage compute ms {timings}")
    pattern = next(s.result for s in run.stages if s.stage == "wafer_pattern")
    assert isinstance(pattern, WaferPatternResult)
    rng = Random(2026)  # one Random: the units first, then the lot
    real_sensor_engine().sample(rng)
    lot = real_wafer_engine().sample(rng).items
    assert [w.source_wafer for w in pattern.wafers] == [w.source_wafer for w in lot]
    final = next(s.result for s in run.stages if s.stage == "final_test")
    assert final is not None
    again, _ = await run_to_end(api, {"source": "sample", "seed": 2026})
    assert again.origin == run.origin
    assert [s.result for s in again.stages] == [s.result for s in run.stages]  # live, repeatable


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
async def test_the_wafer_example_uploads_back_and_runs(api: httpx.AsyncClient, fmt: str) -> None:
    example = await api.get(f"/api/v1/uploads/examples/wafers.{fmt}")
    assert example.status_code == 200
    assert "Wu, M.-J." in example.headers["x-data-attribution"]
    parsed = WaferDieFile().parse(example.content, f"example.{fmt}")
    assert parsed.issues == []
    picked = real_wafer_engine().sample(Random(0)).items
    assert [(w.lot, w.wafer) for w in parsed.items] == [(w.lot, w.wafer) for w in picked]
    assert [w.map for w in parsed.items] == [w.map for w in picked]  # every die, exactly
    if fmt == "xlsx":
        sheets = load_workbook(io.BytesIO(example.content)).sheetnames
        assert sheets == ["Dies", "About"]
    sent = await api.post("/api/v1/uploads", files={"file": (f"lot.{fmt}", example.content)})
    assert sent.status_code == 201
    summary = sent.json()
    assert summary["kind"] == "wafers"
    assert summary["runnable"] is True
    assert summary["wafers"] == len(picked)
    assert summary["lots"] == 1
    run, events = await run_to_end(api, {"source": "upload", "upload_id": summary["upload_id"]})
    assert run.status == "done"
    ends = closing(events)
    assert ends["lot"] == "stage_result"
    assert ends["sensor_risk"] == "stage_skipped"
    assert ends["final_test"] == "stage_skipped"  # a die file carries no recorded labels


async def test_a_drawing_on_a_real_layout_runs_and_one_wafer_skips_the_lot_test(
    api: httpx.AsyncClient,
) -> None:
    layouts = (await api.get("/api/v1/runs/layouts")).json()["items"]
    assert len(layouts) == 6
    layout = layouts[0]
    assert layout["source_wafer"] in layout["rule"]
    grid = decode_map(WaferMap.model_validate(layout["map"])).copy()
    rows, cols = grid.shape
    for c in range(cols):  # a horizontal scratch across the middle row, on dies only
        if grid[rows // 2, c] == 1:
            grid[rows // 2, c] = 2
    drawn = {
        "lot": "mine",
        "wafer": 1,
        "map": {"height": rows, "width": cols, "cells": base64.b64encode(grid.tobytes()).decode()},
        "source_wafer": layout["source_wafer"],
    }
    run, events = await run_to_end(api, {"source": "drawn", "wafers": [drawn]})
    assert run.origin.title == "Your drawing"
    skipped = {e.stage: e.needs for e in events if e.type == "stage_skipped"}
    assert skipped["lot"] == NEEDS_LOT
    pattern = next(s.result for s in run.stages if s.stage == "wafer_pattern")
    assert isinstance(pattern, WaferPatternResult)
    assert pattern.wafers[0].source_wafer is None  # a drawing is not a WM-811K wafer


async def test_the_case_catalog_is_computed_from_the_data(api: httpx.AsyncClient) -> None:
    deadline = time.monotonic() + 600
    while True:
        answer = await api.get("/api/v1/runs/cases")
        if answer.status_code == 200:
            break
        assert answer.status_code == 503
        assert answer.headers["retry-after"] == "5"
        assert time.monotonic() < deadline
        await asyncio.sleep(1)
    cases = {c["case_id"]: c for c in answer.json()["items"]}
    assert set(cases) == {"spe_alarm", "slow_drift", "scratch", "unknown_pattern", "lot_repeat"}
    run, _ = await run_to_end(api, {"source": "case", "case_id": "scratch"})
    pattern = next(s.result for s in run.stages if s.stage == "wafer_pattern")
    assert isinstance(pattern, WaferPatternResult)
    assert pattern.wafers[0].top_label == "Scratch"
    assert pattern.wafers[0].auto_accepted
    unknown, _ = await run_to_end(api, {"source": "case", "case_id": "unknown_pattern"})
    found = next(s.result for s in unknown.stages if s.stage == "wafer_pattern")
    assert isinstance(found, WaferPatternResult)
    assert found.wafers[0].is_unknown
    lot, _ = await run_to_end(api, {"source": "case", "case_id": "lot_repeat"})
    tested = next(s.result for s in lot.stages if s.stage == "lot")
    assert isinstance(tested, LotResult)
    assert tested.lots[0].flagged
    assert any(a.kind == "lot_excursion" for a in lot.alarms)
